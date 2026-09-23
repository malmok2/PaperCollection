"""Excel export of the cumulative DB (replaces the Codex-only @oai/artifact-tool script)."""
from collections import Counter
from datetime import date

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

NAVY, PALE_BLUE, PALE_GRAY, BORDER = "17324D", "EAF2F8", "F4F6F7", "D5DCE3"
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
HEADER_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color=BORDER)


def as_date(value):
    try:
        return date.fromisoformat((value or "")[:10])
    except ValueError:
        return value or None


def data_sheet(wb, title, headers, rows, widths, table_name, date_cols=(), wrap_cols=()):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 32
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    last = max(len(rows) + 1, 2)
    for col in date_cols:
        for (cell,) in ws.iter_rows(min_row=2, max_row=last, min_col=col, max_col=col):
            cell.number_format = "yyyy-mm-dd"
    for col in wrap_cols:
        for (cell,) in ws.iter_rows(min_row=2, max_row=last, min_col=col, max_col=col):
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    if rows:
        table = Table(displayName=table_name, ref=f"A1:{get_column_letter(len(headers))}{last}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
        ws.add_table(table)
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    return ws


def small_table(ws, top, left, headers, rows):
    for j, header in enumerate(headers):
        cell = ws.cell(row=top, column=left + j, value=header)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    for i, row in enumerate(rows, 1):
        for j, value in enumerate(row):
            cell = ws.cell(row=top + i, column=left + j, value=value)
            cell.border = Border(bottom=THIN)
    return top + len(rows)


def build(payload, output_path):
    papers, runs = payload["papers"], payload["runs"]
    wb = Workbook()
    dash = wb.active
    dash.title = "Dashboard"

    data_sheet(
        wb, "Papers",
        ["DOI", "Title EN", "Title KO", "Journal", "ISSN", "Authors", "DOI Registered", "Published Online", "Issue Date",
         "Document Type", "Primary Topic", "Countries", "Primary Country", "Abstract", "Abstract Source", "Abstract URL",
         "DOI URL", "Open Access", "First Seen", "Last Checked", "Status", "Source"],
        [[p["doi"], p["title_en"], p.get("title_ko") or "", p["journal"], p.get("issn") or "", p.get("authors") or "",
          as_date(p.get("doi_registered_date")), as_date(p.get("published_online")), as_date(p.get("issue_date")),
          p["document_type"], p["primary_topic"], p.get("countries") or "", p.get("primary_country") or "",
          p.get("abstract") or "", p.get("abstract_source") or "", p.get("abstract_url") or "", p["doi_url"],
          bool(p.get("open_access")), as_date(p.get("first_seen")), p.get("last_checked"), p.get("status"), p.get("source")]
         for p in papers],
        [30, 58, 45, 34, 12, 42, 14, 14, 14, 14, 18, 14, 10, 65, 12, 40, 38, 11, 13, 13, 10, 22],
        "PapersTable", date_cols=(7, 8, 9, 19), wrap_cols=(2, 3, 6),
    )
    data_sheet(
        wb, "Collection Runs",
        ["Run ID", "Period Start", "Period End", "Executed At", "Found", "New", "Updated", "Errors", "Source Note"],
        [[r["run_id"], as_date(r["period_start"]), as_date(r["period_end"]), r["executed_at"], r["found_count"],
          r["new_count"], r["updated_count"], r["error_count"], r["source_note"]] for r in runs],
        [14, 13, 13, 22, 10, 10, 10, 10, 50], "CollectionRunsTable", date_cols=(2, 3),
    )
    data_sheet(wb, "Topics", ["Topic ID", "Topic KO", "Keywords", "Active"],
               [[t["topic_id"], t["topic_ko"], t["keywords"], bool(t["active"])] for t in payload["topics"]],
               [10, 22, 90, 10], "TopicsTable", wrap_cols=(3,))
    data_sheet(wb, "Paper Topics", ["DOI", "Topic KO", "Primary"],
               [[t["doi"], t["topic_ko"], bool(t["is_primary"])] for t in payload["paper_topics"]],
               [34, 22, 10], "PaperTopicsTable")

    # Dashboard: summary tables + native Excel charts.
    latest = runs[0]["run_id"] if runs else ""
    dash.sheet_view.showGridLines = False
    dash.merge_cells("A1:H2")
    dash["A1"] = "Nuclear Literature Database"
    dash["A1"].font = Font(bold=True, color="FFFFFF", size=20)
    dash["A1"].fill = HEADER_FILL
    dash["A1"].alignment = Alignment(vertical="center")
    dash.merge_cells("A3:H3")
    dash["A3"] = f"{len({p['journal'] for p in papers})} journals · DOI-registered publications · Updated {latest}"
    dash["A3"].font = Font(italic=True, color=NAVY)
    dash["A3"].fill = PatternFill("solid", fgColor=PALE_BLUE)

    stats = [("Total papers", len(papers)), ("Collection runs", len(runs)),
             ("Research articles", sum(p["document_type"] == "연구논문" for p in papers)),
             ("Open access", sum(bool(p.get("open_access")) for p in papers)),
             ("With country", sum(bool(p.get("countries")) for p in papers)),
             ("With abstract", sum(bool(p.get("abstract")) for p in papers))]
    for i, (label, value) in enumerate(stats, 5):
        dash.cell(row=i, column=1, value=label).font = Font(bold=True, color=NAVY)
        cell = dash.cell(row=i, column=2, value=value)
        cell.font = Font(bold=True, color="2166AC", size=14)
        for col in (1, 2):
            dash.cell(row=i, column=col).fill = PatternFill("solid", fgColor=PALE_GRAY)

    types = Counter(p["document_type"] for p in papers).most_common()
    small_table(dash, 5, 4, ["Document type", "Count"], types)

    journals = Counter(p["journal"] for p in papers).most_common()
    j_end = small_table(dash, 13, 1, ["Journal", "Papers"], journals)
    topics = Counter(p["primary_topic"] for p in papers).most_common()
    t_end = small_table(dash, 13, 4, ["Primary topic", "Papers"], topics)
    run_rows = [(r["period_end"], r["found_count"]) for r in reversed(runs)]
    r_end = small_table(dash, 13, 7, ["Period end", "Found"], run_rows)

    for title, col, end, anchor, kind in [("Papers by journal", 1, j_end, "J2", BarChart),
                                          ("Papers by primary topic", 4, t_end, "J20", BarChart),
                                          ("Papers collected per run", 7, r_end, "A30", LineChart)]:
        chart = kind()
        chart.title, chart.legend, chart.height, chart.width = title, None, 8, 16
        if kind is BarChart:
            chart.type = "bar"
        chart.add_data(Reference(dash, min_col=col + 1, min_row=13, max_row=max(end, 14)), titles_from_data=True)
        chart.set_categories(Reference(dash, min_col=col, min_row=14, max_row=max(end, 14)))
        dash.add_chart(chart, anchor)

    for col, width in zip("ABCDEFGH", [34, 12, 3, 24, 10, 3, 14, 10]):
        dash.column_dimensions[col].width = width
    note_row = max(j_end, t_end, r_end) + 16
    dash.cell(row=note_row, column=1, value=(
        "Data source: Crossref DOI created date, OpenAlex (countries, open access, abstracts). DOI is the unique key. "
        "This file is regenerated weekly from data/literature.sqlite; manual edits will be overwritten."
    )).font = Font(italic=True, color=NAVY)
    dash.freeze_panes = "A4"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path
